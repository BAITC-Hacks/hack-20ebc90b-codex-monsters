"""B's public boundary consumes A's actual models, without alternate schemas."""
import json
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pytest

from ekt.contracts import DomainError, MappingConfig, SnapshotManifest, SourceManifest, load_table
from ekt.data import build_snapshot
from ekt.data.artifacts import checksum


AS_OF = datetime.fromisoformat("2026-09-22T23:59:59+05:00")


def master():
    return {"sku_id": "00123", "name": "Contract compatibility fixture", "base_uom": "piece",
            "supplier_id": "test-supplier", "purchase_uom": "pack",
            "base_units_per_purchase_uom": Decimal("10"), "quantity_quantum": Decimal("1")}


def test_real_source_and_mapping_models_return_typed_snapshot(tmp_path):
    source = SourceManifest(
        as_of=AS_OF, mode="synthetic_demo", output_root=str(tmp_path),
        sources=[{"source_id": "master", "kind": "sku_master", "provenance": "synthetic", "rows": [master()]}],
    )
    mapping = MappingConfig(mapping_version="runtime-v1", options={"sources": {"master": {}}})
    result = build_snapshot(source, mapping)
    assert isinstance(result, SnapshotManifest)
    assert result.as_of == AS_OF
    assert result.mapping_version == "runtime-v1"
    assert Path(result.tables["sku_master"].uri).is_relative_to(tmp_path)
    rows = load_table(result, "sku_master")
    assert rows[0]["sku_id"] == "00123"
    assert rows[0]["base_units_per_purchase_uom"] == Decimal("10")
    assert source.output_root == str(tmp_path)  # Boundary does not mutate callers.


def test_old_dict_configuration_is_still_accepted_with_typed_result(tmp_path):
    result = build_snapshot(
        {"as_of": AS_OF.isoformat(), "mode": "synthetic_demo", "artifact_root": str(tmp_path),
         "sources": [{"source_id": "master", "kind": "sku_master", "rows": [master()], "provenance": "synthetic"}]},
        {"version": "old-v1", "sources": {}},
    )
    assert isinstance(result, SnapshotManifest)
    assert result.mapping_version == "old-v1"


def test_mapping_columns_and_output_option_apply_to_registered_sources(tmp_path):
    path = tmp_path / "source.json"
    row = master()
    row["SKU"] = row.pop("sku_id")
    path.write_text(json.dumps([row], default=str), encoding="utf-8")
    source = SourceManifest(
        source_ids=["registered"], as_of=AS_OF, mode="synthetic_demo",
        source_refs=[{"source_id": "registered", "kind": "sku_master", "local_ref": str(path), "checksum": checksum(path)}],
    )
    columns = {key: key for key in master()}
    columns["sku_id"] = "SKU"
    mapping = MappingConfig(mapping_version="registered-v1", columns={"registered": columns}, options={
        "output_root": str(tmp_path / "output"),
        "source_options": {"registered": {"provenance": "synthetic"}},
        "sources": {"registered": {}},
    })
    result = build_snapshot(source, mapping)
    assert load_table(result, "sku_master")[0]["sku_id"] == "00123"
    assert Path(result.tables["sku_master"].uri).is_relative_to(tmp_path / "output")
    assert result.source_refs[0].checksum == checksum(path)


def test_stale_registered_checksum_becomes_shared_domain_error(tmp_path):
    path = tmp_path / "source.json"
    path.write_text(json.dumps([master()], default=str))
    source = SourceManifest(as_of=AS_OF, source_refs=[
        {"source_id": "master", "kind": "sku_master", "local_ref": str(path), "checksum": "stale"}])
    with pytest.raises(DomainError) as caught:
        build_snapshot(source, MappingConfig(options={"output_root": str(tmp_path / "output")}))
    assert caught.value.code == "SOURCE_CHECKSUM_MISMATCH"
    assert not (tmp_path / "output").exists()


def test_conflicting_column_sources_fail_explicitly(tmp_path):
    source = SourceManifest(as_of=AS_OF, output_root=str(tmp_path), sources=[
        {"source_id": "master", "kind": "sku_master", "rows": [master()]}])
    mapping = MappingConfig(columns={"master": {"sku_id": "SKU"}}, options={
        "sources": {"master": {"columns": {"sku_id": "different"}}}})
    with pytest.raises(DomainError) as caught:
        build_snapshot(source, mapping)
    assert caught.value.code == "CONFLICTING_COLUMN_MAPPING"


def registered_systeme_sources(tmp_path):
    from openpyxl import Workbook

    def write(name, headers, rows):
        path = tmp_path / name
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Лист_1"
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        return path

    sku_ids = [f"{number:05}" for number in range(1, 22)]
    sales = write("registered-sales.xlsx", ["Дата", "Номер", "Документ", "Код", "Номенклатура", "Ед.", "Склад", "Количество"], [
        ["01.09.2026 12:00:00", f"doc-{sku}", f"Расходная накладная doc-{sku}", sku, "Registry fixture", "шт", "Алматы", 2]
        for sku in sku_ids])
    terms = write("registered-terms.xlsx", ["№", "Номенклатура", "Номенклатура.Код", "Артикул", "Кратность"], [
        [index, "Registry fixture", sku, f"article-{sku}", 5] for index, sku in enumerate(sku_ids)])
    stock_candidate = write("registered-candidate.xlsx", ["unverified"], [["candidate metadata"]])
    return [{"source_id": source_id, "kind": kind, "local_ref": str(path), "checksum": checksum(path)}
            for source_id, kind, path in [("S07", "supplier_terms", terms), ("S08", "sales_events", sales),
                                          ("S12", "inventory_snapshots", stock_candidate)]]


def test_named_preset_consumes_only_registered_paths_and_selects_twenty(tmp_path, monkeypatch):
    refs = registered_systeme_sources(tmp_path)
    from ekt.data import presets
    def prohibit_discovery(*_args):
        raise AssertionError("Registry adapter must not discover unregistered files")
    monkeypatch.setattr(presets, "_single", prohibit_discovery)
    source = SourceManifest(source_ids=[ref["source_id"] for ref in refs], source_refs=refs, as_of=AS_OF)
    result = build_snapshot(source, MappingConfig(mapping_version="systeme-preview-v1", options={"output_root": str(tmp_path / "output")}))
    assert isinstance(result, SnapshotManifest)
    assert result.tables["sku_master"].row_count == 20
    assert result.tables["sales_events"].row_count == 20
    assert result.tables["supplier_terms"].row_count == 20
    assert result.tables["inventory_snapshots"].row_count == 0
    assert result.mode == "real_preview"
    assert not result.quality.capabilities.can_plan
    assert "S01" not in {ref.source_id for ref in result.source_refs}
    assert {ref.local_ref for ref in result.source_refs if not ref.local_ref.startswith("inline:")} == {ref["local_ref"] for ref in refs}
    assert load_table(result, "sku_master")[0]["sku_id"].startswith("000")


def test_named_preset_rejects_incomplete_registered_aliases(tmp_path):
    refs = registered_systeme_sources(tmp_path)[:2]
    source = SourceManifest(source_ids=[ref["source_id"] for ref in refs], source_refs=refs, as_of=AS_OF)
    with pytest.raises(DomainError) as caught:
        build_snapshot(source, MappingConfig(mapping_version="systeme-preview-v1", options={"output_root": str(tmp_path / "output")}))
    assert caught.value.code == "INVALID_PRESET_SOURCES"
    assert not (tmp_path / "output").exists()
