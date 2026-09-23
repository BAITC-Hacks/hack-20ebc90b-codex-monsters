"""Read local registered source files without changing them or guessing headers."""

from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from .artifacts import checksum, digest_payload
from .errors import DataError


def decimal_value(value, field="quantity", nullable=False):
    if value is None or value == "":
        if nullable:
            return None
        raise DataError("MISSING_QUANTITY", f"Missing {field}")
    if isinstance(value, bool):
        raise DataError("INVALID_QUANTITY", f"Invalid {field}")
    try:
        result = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise DataError("INVALID_QUANTITY", f"Invalid {field}") from None
    if not result.is_finite() or len(result.as_tuple().digits) > 38 or result.as_tuple().exponent < -18 or result.adjusted() >= 38:
        raise DataError("INVALID_QUANTITY", f"Non-finite or unsupported precision for {field}")
    return result


def timestamp(value, field="timestamp") -> str:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time(), timezone.utc)
    else:
        try:
            parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except (ValueError, TypeError):
            raise DataError("INVALID_TIMESTAMP", f"Invalid {field}; expected ISO timestamp") from None
    if parsed.tzinfo is None:
        raise DataError("MISSING_TIMEZONE", f"Explicit source timezone required for {field}")
    return parsed.isoformat()


def source_records(source: dict, mapping: dict):
    """Return (records with source row locators, checksum, actual headers).

    Numeric Excel identifiers are accepted only with an explicit all-zero display
    format, which preserves leading zeros. Other numeric SKU values are rejected
    later instead of pretending lost leading zeros are recoverable.
    """
    if "rows" in source:
        rows = source["rows"]
        headers = list(dict.fromkeys(key for row in rows for key in row))
        result = [(str(index + 1), dict(row)) for index, row in enumerate(rows)]
        fingerprint = digest_payload(rows)
    else:
        path = Path(source["path"]).expanduser().resolve()
        if not path.is_file():
            raise DataError("SOURCE_NOT_FOUND", f"Registered source {source['source_id']} is unavailable")
        fingerprint = checksum(path)
        file_format = source.get("format", path.suffix.lstrip(".").lower())
        header_row = int(source.get("header_row", 1))
        if header_row < 1:
            raise DataError("INVALID_HEADER", "header_row must be one-based")
        if file_format == "csv":
            with path.open(encoding=source.get("encoding", "utf-8-sig"), newline="") as stream:
                reader = csv.reader(stream, delimiter=source.get("delimiter", ","))
                for _ in range(header_row - 1):
                    next(reader, None)
                headers = next(reader, [])
                result = []
                for index, values in enumerate(reader, header_row + 1):
                    if len(values) > len(headers):
                        raise DataError("INVALID_ROW_WIDTH", f"Source {source['source_id']} row {index} has extra cells")
                    result.append((str(index), dict(zip(headers, values))))
        elif file_format in {"xlsx", "xlsm"}:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = source.get("sheet")
                if not sheet or sheet not in workbook.sheetnames:
                    raise DataError("INVALID_SHEET", "An existing worksheet name is required")
                reader = workbook[sheet].iter_rows(min_row=header_row)
                headers = [cell.value for cell in next(reader, [])]
                result = []
                sku_header = mapping.get("columns", {}).get("sku_id", "sku_id")
                for index, cells in enumerate(reader, header_row + 1):
                    values = []
                    for header, cell in zip(headers, cells):
                        value = cell.value
                        fmt = cell.number_format
                        if header == sku_header and isinstance(value, (int, float)) and not isinstance(value, bool):
                            if fmt and set(fmt) == {"0"} and int(value) == value:
                                value = str(int(value)).zfill(len(fmt))
                        values.append(value)
                    result.append((f"{sheet}!{index}", dict(zip(headers, values))))
            finally:
                workbook.close()
        elif file_format == "json":
            rows = json.loads(path.read_text(encoding="utf-8"))
            if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
                raise DataError("INVALID_SOURCE", "JSON source must contain an array of objects")
            headers = list(dict.fromkeys(key for row in rows for key in row))
            result = [(str(index + 1), row) for index, row in enumerate(rows)]
        elif file_format == "parquet":
            import pyarrow.parquet as pq
            table = pq.read_table(path)
            headers = table.column_names
            result = [(str(index + 1), row) for index, row in enumerate(table.to_pylist())]
        else:
            raise DataError("UNSUPPORTED_SOURCE_FORMAT", f"Unsupported source format: {file_format}")
        if checksum(path) != fingerprint:
            raise DataError("SOURCE_CHANGED_DURING_READ", "Source changed while being read; retry on a completed export", retryable=True)
    if any(not isinstance(header, str) or not header for header in headers) or len(headers) != len(set(headers)):
        raise DataError("INVALID_HEADER", f"Source {source['source_id']} has blank or duplicate headers")
    required = set(mapping.get("required_headers", [])) | set(mapping.get("columns", {}).values())
    missing = required - set(headers)
    if missing:
        raise DataError("SOURCE_HEADER_MISMATCH", f"Source {source['source_id']} missing headers: {', '.join(sorted(missing))}")
    return result, fingerprint, headers


def profile_source(source: dict, mapping: dict) -> dict:
    """Return structural evidence, never commercial source rows or identifiers."""
    if source.get("metadata_only"):
        return {"source_id": source["source_id"], "checksum": checksum(source["path"]),
                "metadata_only": True, "row_count": None, "note": "No verified row mapping for this source"}
    records, fingerprint, headers = source_records(source, mapping)
    columns = mapping.get("columns", {})
    quantity_field = columns.get("quantity_base", columns.get("quantity", "quantity_base"))
    date_field = columns.get("event_at", "event_at")
    uom_field = columns.get("base_uom", "base_uom")
    signs = Counter()
    dates = []
    uoms = Counter()
    invalid = 0
    for _, row in records:
        value = row.get(quantity_field)
        if value is not None and value != "":
            try:
                quantity = decimal_value(value)
                signs["positive" if quantity > 0 else "negative" if quantity < 0 else "zero"] += 1
            except DataError:
                invalid += 1
        value = row.get(date_field)
        if value:
            try:
                if isinstance(value, str) and mapping.get("date_formats", {}).get("event_at"):
                    value = datetime.strptime(value, mapping["date_formats"]["event_at"])
                if isinstance(value, datetime) and value.tzinfo is None and mapping.get("timezone"):
                    from zoneinfo import ZoneInfo
                    value = value.replace(tzinfo=ZoneInfo(mapping["timezone"]))
                dates.append(timestamp(value))
            except (DataError, ValueError):
                pass
        if row.get(uom_field):
            uoms[str(row[uom_field])] += 1
    return {"source_id": source["source_id"], "checksum": fingerprint, "headers": headers,
            "row_count": len(records), "quantity_sign_counts": dict(signs), "invalid_quantity_count": invalid,
            "date_min": min(dates) if dates else None, "date_max": max(dates) if dates else None,
            "uom_counts": dict(uoms)}
